"""Maintain the production-shaped Excel source used by the audit-report replay.

The report runtime never imports this module.  It reads the resulting workbook through the
mock HTTP adapters and the financial-data Excel adapter, exactly as it would read independent
production systems.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


PACKAGE_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = (
    PACKAGE_ROOT
    / "fixtures"
    / "audit-report"
    / "source-simulation"
    / "模拟审计系统全量数据.xlsx"
)
ORGANIZATION_ID = "ORG-QF-001"
PROJECT_ID = "PRJ-QF-2025"
PREVIOUS_PROJECT_ID = "PRJ-QF-2023"


def headers(ws) -> dict[str, int]:
    return {str(cell.value): index for index, cell in enumerate(ws[1], start=1)}


def ensure_columns(ws, columns: list[str]) -> dict[str, int]:
    current = headers(ws)
    for column in columns:
        if column not in current:
            ws.cell(row=1, column=ws.max_column + 1, value=column)
            current[column] = ws.max_column
    return current


def replace_sheet(workbook, name: str, rows: list[dict[str, object]]) -> None:
    index = workbook.sheetnames.index(name) if name in workbook.sheetnames else len(workbook.sheetnames)
    if name in workbook.sheetnames:
        del workbook[name]
    ws = workbook.create_sheet(name, index)
    columns = list(dict.fromkeys(field for record in rows for field in record))
    ws.append(columns)
    for record in rows:
        ws.append([record.get(column) for column in columns])


def update_by_key(ws, key: str, values: dict[str, dict[str, object]]) -> None:
    columns = ensure_columns(ws, sorted({field for record in values.values() for field in record}))
    for row in ws.iter_rows(min_row=2):
        identifier = str(row[columns[key] - 1].value or "")
        if identifier not in values:
            continue
        for field, value in values[identifier].items():
            ws.cell(row=row[0].row, column=columns[field], value=value)


def deduplicate_projects(workbook) -> None:
    ws = workbook["审计项目"]
    columns = headers(ws)
    records: list[dict[str, object]] = []
    seen: set[str] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = {cell.value: row[index - 1] for index, cell in enumerate(ws[1], start=1)}
        task_id = str(record.get("taskId") or "")
        if task_id in seen:
            continue
        seen.add(task_id)
        records.append(record)
    if "TASK-QF-REG-PREV" not in seen:
        records.append(
            {
                "caseId": "QIFAN-PREVIOUS-REGULAR",
                "description": "启帆路上一次审计项目，用于离任报告两期问题比对。",
                "taskId": "TASK-QF-REG-PREV",
                "projectId": PREVIOUS_PROJECT_ID,
                "reportType": "regular",
                "organizationId": ORGANIZATION_ID,
                "auditStart": "2021-01-01",
                "auditEnd": "2023-06-30",
                "auditGroupEstablishedMonth": "2023年7月",
                "reportDate": "2023年8月1日",
                "templateId": "regular-branch-report",
                "templateVersion": "2023.06",
                "feedbackCompleted": True,
            }
        )
    replace_sheet(workbook, "审计项目", records)


def update_appointments(workbook) -> None:
    ws = workbook["OA任免发文"]
    update_by_key(
        ws,
        "documentNumber",
        {
            "东证人字[2021]278号": {
                "action": "appoint",
                "fullTitle": "上海浦东新区启帆路证券营业部副总经理（主持工作）",
            },
            "财富委字[2025]270号": {
                "action": "remove",
                "fullTitle": "上海浦东新区启帆路证券营业部副总经理（主持工作）",
            },
            "东证人字[2025]代职01号": {
                "action": "acting",
                "title": "营业部负责人职务，任上海第一分公司（筹）负责人",
                "fullTitle": "营业部负责人职务，任上海第一分公司（筹）负责人",
            },
        },
    )


CURRENT_FINDINGS = {
    "F-001": {
        "policyBasis": "公司《投资者适当性管理办法》规定，分支机构现场告知适当性评估结果等全过程应录音或录像，并妥善保管影像资料。",
        "factText": "抽查21笔客户账户开立业务，其中1笔机构代理人在账户开立的文件签署视频中全程佩戴口罩，无法辨别客户脸部特征。",
        "internalSubitems": "",
    },
    "F-002": {
        "policyBasis": "公司《投资者适当性管理办法》规定，公司向普通投资者提供的交易服务的风险等级应当与投资者分类结果相匹配。公司认为投资者参与相关交易不适当的，应当拒绝提供相关服务。",
        "factText": "审计期内，营业部2名客户适当性管理不到位，已获得的交易服务与其分类结果不匹配，营业部未规范落实后续管理要求。",
        "internalSubitems": "",
    },
    "F-003": {
        "policyBasis": "公司《分支机构反洗钱工作管理办法》规定，分支机构反洗钱工作小组建立议事机制，每年至少召开四次反洗钱工作会议，议事范围含括：审议分支机构反洗钱工作小组成员变动、分支机构反洗钱年度培训和宣传工作计划及分支机构其他重要反洗钱工作事项等。",
        "factText": "审计发现，营业部反洗钱工作小组议事机制执行不规范。",
        "internalSubitems": "（1）会议频次不足。2023-2024年期间，反洗钱工作小组仅召开了一次会议，不符合公司制度的频次要求。|（2）议事范围执行不到位。营业部反洗钱年度培训及宣传工作计划等重要事项，均未按规定履行审议程序。|（3）小组人员管理不及时。营业部财务人员因办公地址变更而无法实际履行反洗钱工作职责，营业部未及时对小组成员名单进行更新调整。",
    },
    "F-004": {
        "policyBasis": "公司《洗钱风险评估及客户分类管理办法》规定，对于新建立业务关系的客户，各单位应在建立业务关系后的10个工作日内划分其风险等级。",
        "factText": "审计期内，营业部个别新开户客户的反洗钱风险等级划分未在10个工作日内完成。",
        "internalSubitems": "",
    },
    "F-005": {
        "policyBasis": "公司《反洗钱操作指引》规定，营业部应在收到风险等级定期审核流程的10个工作日内完成，高风险客户至少每半年审核一次。",
        "factText": "审计发现，营业部部分高风险客户定期审核流程未在规定时间内提交，导致高风险客户定期审核超期。",
        "internalSubitems": "",
    },
    "F-006": {
        "policyBasis": "公司《经纪业务佣金管理办法》规定，客户产生的业务创收应覆盖其交易成本，除个性化佣金费率等特殊情况外，客户佣金标准不得低于最低佣金。",
        "factText": "审计发现，2名已开通可转债交易权限的客户，可转债佣金费率低于公司成本。",
        "internalSubitems": "",
    },
    "F-007": {
        "policyBasis": "公司《证券经纪业务信息公示管理细则》规定，分支机构营业场所应通过分支机构营业场所等载体，公示包括展业人员信息、相关业务信息、费用与佣金信息等。",
        "factText": "审计发现，营业部未按要求公示两融业务信息（融资融券专员、标的证券名单、融资利率及融券费率）及股票期权业务信息（佣金收取标准、保证金收取标准）。",
        "internalSubitems": "",
    },
    "F-008": {
        "policyBasis": "公司《安全保卫和消防安全工作管理办法》规定，安防、消防设施需聘请有专业资质的维护保养单位进行定期维护保养，确保设施正常运转。",
        "factText": "审计发现，营业部未聘请专业的消防维保单位对安防、消防设施进行定期维保。",
        "internalSubitems": "",
    },
    "F-009": {
        "policyBasis": "根据计划财务管理总部关于应付账款清理的相关要求，对于长期挂账应付账款（入账时间五年以上），经确定不再需要支付或认定为确实无法支付的，由营业部申请，经审批后进行清理。",
        "factText": "审计发现，营业部存在1笔账龄7年的项目质保金尾款，营业部未按要求对是否需要清理进行认定。",
        "internalSubitems": "",
    },
}


PREVIOUS_TITLES = [
    "账户业务资料不完备",
    "反洗钱工作小组议事机制执行不规范",
    "反洗钱工作不及时",
    "佣金设置错误",
    "未按要求报备制度",
    "员工未申报持股企业",
    "客户回访及通知不规范",
    "档案管理工作薄弱",
    "信息公示不规范",
    "监控及安防管理存缺陷",
    "系统运维管理不符合要求",
    "费用报销不规范",
]


def update_findings(workbook) -> None:
    ws = workbook["审计发现"]
    columns = ensure_columns(
        ws,
        ["rawDetail", "internalSubitems", "responsibility", "majorType", "majorConfirmed", "sourceOrder"],
    )
    current_rows: list[dict[str, object]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = {name: row[index - 1] for name, index in columns.items()}
        if record.get("projectId") != PROJECT_ID:
            continue
        finding_id = str(record["findingId"])
        update = CURRENT_FINDINGS[finding_id]
        record.update(update)
        record["rawDetail"] = (
            f"问题标题：{record['title']}\n制度依据：{update['policyBasis']}\n审计发现：{update['factText']}"
        )
        record["responsibility"] = "管理责任"
        record["majorType"] = ""
        record["majorConfirmed"] = False
        record["sourceOrder"] = int(finding_id.split("-")[-1])
        current_rows.append(record)

    previous_rows: list[dict[str, object]] = []
    for index, title in enumerate(PREVIOUS_TITLES, start=1):
        finding_id = f"F-PREV-{index:03d}"
        category = "反洗钱工作" if "反洗钱" in title else "综合管理"
        policy_basis = "上一次审计底稿所载制度依据。"
        fact_text = f"上一次审计发现，营业部存在“{title}”问题。"
        subcategory = title
        if title == "反洗钱工作小组议事机制执行不规范":
            policy_basis = CURRENT_FINDINGS["F-003"]["policyBasis"]
            fact_text = "审计发现，营业部反洗钱工作小组议事机制执行不规范。"
            subcategory = "反洗钱工作小组议事机制"
        elif title == "信息公示不规范":
            policy_basis = "公司从业人员信息公示要求规定，营业部应及时更新展业人员信息。"
            fact_text = "上一次审计发现，营业部未及时更新从业人员变更信息。"
            subcategory = "从业人员信息公示"
        previous_rows.append(
            {
                "findingId": finding_id,
                "projectId": PREVIOUS_PROJECT_ID,
                "organizationId": ORGANIZATION_ID,
                "category": category,
                "subcategory": subcategory,
                "findingType": "制度执行类",
                "severity": "一般",
                "title": title,
                "policyBasis": policy_basis,
                "factText": fact_text,
                "issueCount": 1,
                "foundDate": "2023-07-15",
                "status": "closed",
                "isHistorical": True,
                "isRepeat": title == "反洗钱工作小组议事机制执行不规范",
                "isSubjectResponsible": True,
                "rawDetail": f"问题标题：{title}\n制度依据：{policy_basis}\n审计发现：{fact_text}",
                "internalSubitems": "",
                "responsibility": "管理责任",
                "majorType": "",
                "majorConfirmed": False,
                "sourceOrder": index,
            }
        )
    replace_sheet(workbook, "审计发现", current_rows + previous_rows)


def update_rectifications(workbook) -> None:
    ws = workbook["整改记录"]
    columns = headers(ws)
    current = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = {name: row[index - 1] for name, index in columns.items()}
        if record.get("projectId") == PROJECT_ID:
            current.append(record)
    previous = []
    for index, title in enumerate(PREVIOUS_TITLES, start=1):
        finding_id = f"F-PREV-{index:03d}"
        unresolved = title == "反洗钱工作小组议事机制执行不规范"
        previous.append(
            {
                "findingId": finding_id,
                "organizationId": ORGANIZATION_ID,
                "projectId": PREVIOUS_PROJECT_ID,
                "rectificationId": f"RECT-{finding_id}",
                "status": "overdue" if unresolved else "completed",
                "requirement": f"整改“{title}”问题。",
                "deadline": "2023-12-31",
                "completedAt": None if unresolved else "2023-12-20",
            }
        )
    replace_sheet(workbook, "整改记录", current + previous)


def update_risk_events(workbook) -> None:
    period = {
        "organizationId": ORGANIZATION_ID,
        "auditStart": "2023-09-01",
        "auditEnd": "2025-10-31",
    }
    rows = [
        {**period, "eventId": "RISK-SECURITY", "type": "security-incident", "state": "VERIFIED_NONE"},
        {**period, "eventId": "RISK-EMERGENCY", "type": "major-emergency", "state": "VERIFIED_NONE"},
        {**period, "eventId": "RISK-LAWSUIT", "type": "lawsuit", "state": "VERIFIED_NONE"},
        {**period, "eventId": "RISK-COMPLAINT", "type": "complaint", "state": "VERIFIED_NONE"},
        {**period, "eventId": "RISK-REGULATORY", "type": "regulatory-letter", "state": "VERIFIED_NONE"},
        {
            **period,
            "eventId": "RISK-ACCOUNTABILITY",
            "type": "accountability",
            "state": "VERIFIED_VALUE",
            "description": "2025年3月，营业部相关人员因合规管理及审计整改不到位被予以合规问责。",
            "regularDescription": "2025年3月，营业部因未能勤勉履行合规管理要求，且对内部审计发现的问题整改不力、屡审屡犯，根据公司《内部审计工作制度》及《合规问责实施办法》，营业部负责人及合规与风控专员被予以合规问责。",
            "turnoverDescription": "2025年3月，营业部因未能勤勉履行合规管理要求，对内部审计发现的问题整改不力且屡审屡犯，公司对卢俊同志及其他相关人员进行了合规问责。",
            "occurredAt": "2025-03-15",
        },
    ]
    replace_sheet(workbook, "风险事项", rows)


def update_operating_data(workbook) -> None:
    ws = workbook["经营数据"]
    if ws.cell(1, 4).value != "2022年":
        ws.insert_cols(4, 2)
    ws.cell(1, 4, "2022年")
    ws.cell(1, 5, None)
    ws.cell(2, 4, "金额")
    ws.cell(2, 5, "排名")
    values = {
        "考核收入": (825.16, 86),
        "其中：代理买卖证券业务净收入": (310.53, 79),
        "客户保证金利息净收入": (135.73, None),
        "产品相关收入": (234.80, 81),
        "两融息差净收入": (110.61, 80),
        "其他收入": (33.49, None),
        "业务管理费": (667.72, None),
        "考核利润": (153.16, 83),
        "客户资产规模（万元）": (115190.10, 95),
        "股基交易量（万元）": (1489060.33, 96),
        "非货币产品保有量（万元）": (14466.08, 98),
        "KPI（百分制）": (64.80, 40),
    }
    for row in ws.iter_rows(min_row=3):
        metric = str(row[2].value or "")
        if metric in values:
            row[3].value, row[4].value = values[metric]

    participants = [
        {"period": "2022年", "participants": 166},
        {"period": "2023年", "participants": 168},
        {"period": "2024年", "participants": 169},
        {"period": "2025年1-9月", "participants": 169},
    ]
    replace_sheet(workbook, "排名参与家数", participants)


def update_aml(workbook) -> None:
    update_by_key(
        workbook["反洗钱汇总"],
        "organizationId",
        {
            ORGANIZATION_ID: {
                "suspiciousTransactionCount": 15,
                "suspiciousTransactionType": "一般可疑交易",
                "generalSuspiciousTransactionCount": 15,
                "keySuspiciousTransactionCount": 0,
                "problemQueryComplete": True,
                "majorMatterQueryComplete": True,
            }
        },
    )
    period = {"organizationId": ORGANIZATION_ID, "auditStart": "2023-09-01", "auditEnd": "2025-10-31"}
    replace_sheet(
        workbook,
        "总部可疑交易认定",
        [
            {
                **period,
                "recognitionId": "AML-SUS-001",
                "recognitionResult": "认定",
                "recognitionDate": "2025-09-30",
                "suspiciousType": "一般可疑交易",
                "transactionCount": 15,
                "generalCount": 15,
                "keyCount": 0,
            }
        ],
    )
    update_by_key(
        workbook["反洗钱协查函"],
        "letterId",
        {
            "AML-LETTER-002": {
                "enteredDate": "2025-08-08",
                "reviewedDate": "2025-08-11",
                "overdue": False,
                "riskAdjustmentStatus": "已调整",
                "findingIds": "",
            }
        },
    )


def update_narrative_facts(workbook) -> None:
    update_by_key(
        workbook["审计叙述事实"],
        "organizationId",
        {
            ORGANIZATION_ID: {
                "managerDutySummary": "审计期内，卢俊同志基本能够贯彻执行国家有关证券市场发展的方针政策，遵守《证券法》《证券公司内部控制指引》《证券经纪业务管理办法》等法律法规以及证券业务规则，执行廉洁从业规定，公司未受理或办理过涉及卢俊同志个人的信访及案件。审计期内，卢俊同志基本能够执行公司有关经纪业务的重大经营决策，平稳开展业务。2023-2024年度，公司对卢俊同志的绩效考核结果均为B。",
                "previousRectificationSummary": "“反洗钱工作小组议事机制执行不规范”问题在本次审计中仍然存在，未有效整改。",
            }
        },
    )


def main() -> None:
    workbook = load_workbook(WORKBOOK)
    deduplicate_projects(workbook)
    update_appointments(workbook)
    update_findings(workbook)
    update_rectifications(workbook)
    update_risk_events(workbook)
    update_operating_data(workbook)
    update_aml(workbook)
    update_narrative_facts(workbook)
    workbook.save(WORKBOOK)
    print(WORKBOOK)


if __name__ == "__main__":
    main()
