from __future__ import annotations

import argparse
import json
import re
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from docx import Document


def normalize(text: str) -> str:
    return re.sub(r"[\s，。；：、“”‘’（）()《》\-—]", "", text)


def document_text(path: Path) -> tuple[str, list[str]]:
    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    table_cells = [
        cell.text.strip()
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    ]
    return "\n".join([*paragraphs, *table_cells]), paragraphs


def contains(text: str, value: str) -> bool:
    return normalize(value) in normalize(text)


def number_token(value: float | int) -> str:
    if isinstance(value, int):
        return str(value)
    return f"{value:.8f}".rstrip("0").rstrip(".")


def numeric_coverage(text: str, dataset: dict[str, Any]) -> tuple[int, int, float]:
    if dataset["task"]["reportType"] == "aml":
        return 0, 0, 1.0
    expected = [
        canonical_number(number_token(point["value"]))
        for metric in dataset["operatingMetrics"]
        for point in metric["points"]
    ]
    expected.extend(
        canonical_number(str(point["rank"]))
        for metric in dataset["operatingMetrics"]
        for point in metric["points"]
        if point.get("rank") is not None
    )
    if not expected:
        return 0, 0, 1.0
    actual = {
        canonical_number(token)
        for token in re.findall(r"(?<!\d)-?\d[\d,]*(?:\.\d+)?", text)
    }
    matched = sum(1 for value in expected if value in actual)
    return matched, len(expected), matched / len(expected)


def canonical_number(value: str) -> str:
    normalized = value.replace(",", "")
    if "." not in normalized:
        return normalized
    return normalized.rstrip("0").rstrip(".")


def structural_docx_qa(path: Path, draft: dict[str, Any]) -> tuple[int, int, list[str]]:
    document = Document(str(path))
    expected_tables = [
        table
        for section in draft["sections"]
        for table in [
            *section.get("tables", []),
            *[
                table
                for subsection in section.get("subsections", [])
                for table in subsection.get("tables", [])
            ],
        ]
    ]
    checks: list[tuple[bool, str]] = []
    checks.append(
        (
            len(document.tables) == len(expected_tables),
            f"表格数量：实际{len(document.tables)}，预期{len(expected_tables)}",
        )
    )
    for index, expected in enumerate(expected_tables):
        if index >= len(document.tables):
            break
        actual = document.tables[index]
        expected_columns = len(expected["headers"])
        actual_columns = len(actual.columns)
        checks.append(
            (
                actual_columns == expected_columns,
                f"表{index + 1}列数：实际{actual_columns}，预期{expected_columns}",
            )
        )
        checks.append(
            (
                all(len(row.cells) == expected_columns for row in actual.rows),
                f"表{index + 1}不存在幽灵空白列",
            )
        )
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for section in draft["sections"]:
        for closing in section.get("closingParagraphs", []):
            closing_text = closing["text"]
            closing_index = paragraphs.index(closing_text) if closing_text in paragraphs else -1
            preceding_texts = [
                paragraph["text"]
                for subsection in section.get("subsections", [])
                for paragraph in subsection.get("paragraphs", [])
            ]
            preceding_indexes = [
                paragraphs.index(text) for text in preceding_texts if text in paragraphs
            ]
            checks.append(
                (
                    closing_index >= 0
                    and (not preceding_indexes or closing_index > max(preceding_indexes)),
                    f"{closing['paragraphId']}位于意见子项之后",
                )
            )
    passed = sum(1 for value, _ in checks if value)
    failures = [description for value, description in checks if not value]
    return passed, len(checks), failures


def finding_coverage(text: str, dataset: dict[str, Any]) -> tuple[list[str], list[str], float]:
    matched = [
        finding["findingId"]
        for finding in dataset["findings"]
        if contains(text, finding["title"])
    ]
    missing = [
        finding["findingId"]
        for finding in dataset["findings"]
        if finding["findingId"] not in matched
    ]
    return matched, missing, len(matched) / max(len(dataset["findings"]), 1)


def section_coverage(text: str, draft: dict[str, Any]) -> tuple[list[str], list[str], float]:
    headings = [
        section["heading"]
        for section in draft["sections"]
    ]
    matched = [heading for heading in headings if contains(text, heading)]
    missing = [heading for heading in headings if heading not in matched]
    return matched, missing, len(matched) / max(len(headings), 1)


def key_fact_coverage(text: str, dataset: dict[str, Any]) -> tuple[list[str], list[str], float]:
    key_facts = {
        "机构名称": dataset["organization"]["fullName"],
        "审计开始年份": dataset["task"]["auditStart"][:4],
        "审计结束年份": dataset["task"]["auditEnd"][:4],
    }
    if dataset["task"]["reportType"] == "regular":
        key_facts.update(
            {
                "营业地址": dataset["organization"]["address"],
                "营业面积": str(dataset["organization"]["areaSquareMeters"]),
                "员工人数": str(dataset["personnel"]["employeeCount"]),
            }
        )
    matched = [label for label, value in key_facts.items() if contains(text, value)]
    missing = [label for label in key_facts if label not in matched]
    return matched, missing, len(matched) / max(len(key_facts), 1)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--human", required=True, type=Path)
    parser.add_argument("--generated", required=True, type=Path)
    parser.add_argument("--dataset", required=True, type=Path)
    parser.add_argument("--draft", required=True, type=Path)
    parser.add_argument("--output-json", required=True, type=Path)
    parser.add_argument("--output-md", required=True, type=Path)
    args = parser.parse_args()

    dataset = json.loads(args.dataset.read_text(encoding="utf-8"))
    draft = json.loads(args.draft.read_text(encoding="utf-8"))
    human_text, _ = document_text(args.human)
    generated_text, _ = document_text(args.generated)

    generated_findings, generated_missing_findings, generated_finding_rate = finding_coverage(
        generated_text, dataset
    )
    human_findings, human_missing_findings, human_finding_rate = finding_coverage(
        human_text, dataset
    )
    generated_numeric_count, numeric_total, generated_numeric_rate = numeric_coverage(
        generated_text, dataset
    )
    human_numeric_count, _, human_numeric_rate = numeric_coverage(human_text, dataset)
    generated_sections, generated_missing_sections, generated_section_rate = section_coverage(
        generated_text, draft
    )
    human_sections, human_missing_sections, human_section_rate = section_coverage(
        human_text, draft
    )
    generated_facts, generated_missing_facts, generated_fact_rate = key_fact_coverage(
        generated_text, dataset
    )
    human_facts, human_missing_facts, human_fact_rate = key_fact_coverage(
        human_text, dataset
    )
    text_similarity = SequenceMatcher(
        None, normalize(human_text), normalize(generated_text)
    ).ratio()
    structural_passed, structural_total, structural_failures = structural_docx_qa(
        args.generated, draft
    )
    structural_rate = structural_passed / max(structural_total, 1)
    completion_score = round(
        (
            generated_finding_rate * 0.4
            + generated_numeric_rate * 0.3
            + generated_section_rate * 0.15
            + generated_fact_rate * 0.1
            + structural_rate * 0.05
        )
        * 100,
        2,
    )
    result = {
        "humanReport": str(args.human),
        "generatedReport": str(args.generated),
        "completionScore": completion_score,
        "textSimilarityReferenceOnly": round(text_similarity * 100, 2),
        "generated": {
            "findingCoverage": round(generated_finding_rate * 100, 2),
            "matchedFindingIds": generated_findings,
            "missingFindingIds": generated_missing_findings,
            "numericCoverage": round(generated_numeric_rate * 100, 2),
            "numericMatched": generated_numeric_count,
            "numericExpected": numeric_total,
            "sectionCoverage": round(generated_section_rate * 100, 2),
            "matchedSections": generated_sections,
            "missingSections": generated_missing_sections,
            "keyFactCoverage": round(generated_fact_rate * 100, 2),
            "matchedKeyFacts": generated_facts,
            "missingKeyFacts": generated_missing_facts,
            "structuralDocxQa": round(structural_rate * 100, 2),
            "structuralDocxFailures": structural_failures,
            "visualQa": "需以Word逐页渲染结果单独确认，不再硬编码分数",
        },
        "humanBaseline": {
            "findingCoverage": round(human_finding_rate * 100, 2),
            "matchedFindingIds": human_findings,
            "missingFindingIds": human_missing_findings,
            "numericCoverage": round(human_numeric_rate * 100, 2),
            "numericMatched": human_numeric_count,
            "numericExpected": numeric_total,
            "sectionCoverage": round(human_section_rate * 100, 2),
            "matchedSections": human_sections,
            "missingSections": human_missing_sections,
            "keyFactCoverage": round(human_fact_rate * 100, 2),
            "matchedKeyFacts": human_facts,
            "missingKeyFacts": human_missing_facts,
        },
        "interpretation": {
            "completionScore": "按问题40%、数值30%、主章节15%、关键事实10%、DOCX结构QA 5%计算。",
            "textSimilarity": "仅作为措辞接近度参考，不纳入完成度；审计报告允许在不改变事实的前提下改写。",
            "visualQa": "视觉版式必须基于Word渲染逐页检查，自动对比不再假定为100分。",
        },
    }
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_md.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    args.output_md.write_text(
        f"""# 智能体报告与人工报告对比

## 总体结论

- 加权完成度：{completion_score:.2f} 分
- 报告问题覆盖：{generated_finding_rate * 100:.2f}%（人工基线 {human_finding_rate * 100:.2f}%）
- 经营数值与排名覆盖：{generated_numeric_rate * 100:.2f}%（人工基线 {human_numeric_rate * 100:.2f}%）
- 主章节覆盖：{generated_section_rate * 100:.2f}%（人工基线 {human_section_rate * 100:.2f}%）
- 关键事实覆盖：{generated_fact_rate * 100:.2f}%（人工基线 {human_fact_rate * 100:.2f}%）
- DOCX 结构 QA：{structural_rate * 100:.2f}%（{structural_passed}/{structural_total}）
- Word 逐页视觉 QA：需单独渲染确认
- 全文字符序列相似度：{text_similarity * 100:.2f}%（仅作为措辞接近度参考，不计入完成度）

## 缺失项

- 智能体缺失问题：{"、".join(generated_missing_findings) or "无"}
- 智能体缺失章节：{"、".join(generated_missing_sections) or "无"}
- 智能体缺失关键事实：{"、".join(generated_missing_facts) or "无"}
- DOCX 结构异常：{"；".join(structural_failures) or "无"}

## 评分说明

完成度按问题覆盖 40%、数值与排名覆盖 30%、主章节覆盖 15%、关键事实覆盖 10%、DOCX 结构 QA 5% 计算。视觉版式必须在 Word 渲染后逐页确认，不再由脚本硬编码为 100 分。文本相似度不纳入完成度，因为报告允许在不改变事实、数量、对象、期间和制度依据的前提下进行语言优化。

## 边界

本次使用人工报告回放事实构造模拟输入，证明生成链路能够恢复人工报告的主要事实和结构；它不是独立新项目的盲测。真实完成度仍需使用未参与规则编制的新营业部数据做留出集验证。
""",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
