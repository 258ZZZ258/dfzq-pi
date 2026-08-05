from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from docx import Document


REPORT_TYPES = {
    "反洗钱": "aml",
    "离任": "turnover",
    "常规": "regular",
}


def compact(value: str) -> str:
    return re.sub(r"\s+", "", value)


def report_type(name: str) -> str | None:
    for marker, value in REPORT_TYPES.items():
        if marker in name:
            return value
    return None


def block_text(document: Any) -> list[dict[str, Any]]:
    body = document.element.body
    paragraph_by_element = {paragraph._p: paragraph for paragraph in document.paragraphs}
    table_by_element = {table._tbl: table for table in document.tables}
    blocks: list[dict[str, Any]] = []
    for element in body.iterchildren():
        if element in paragraph_by_element:
            paragraph = paragraph_by_element[element]
            text = paragraph.text.strip()
            if text:
                blocks.append(
                    {
                        "kind": "paragraph",
                        "text": text,
                        "style": paragraph.style.name if paragraph.style else "",
                    }
                )
        elif element in table_by_element:
            table = table_by_element[element]
            rows = [
                [compact(cell.text) for cell in row.cells]
                for row in table.rows
            ]
            blocks.append({"kind": "table", "rows": rows})
    return blocks


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    inventory = []
    # Office creates lock files such as "~$报告.docx" while a document is open.
    # They are not valid DOCX packages and must never enter the ingestion set.
    for path in sorted(path for path in args.input_dir.glob("*.docx") if not path.name.startswith("~$")):
        kind = report_type(path.name)
        if kind is None:
            continue
        document = Document(str(path))
        blocks = block_text(document)
        paragraphs = [block for block in blocks if block["kind"] == "paragraph"]
        tables = [block for block in blocks if block["kind"] == "table"]
        inventory.append(
            {
                "fileName": path.name,
                "path": str(path),
                "reportType": kind,
                "paragraphCount": len(paragraphs),
                "tableCount": len(tables),
                "blocks": blocks,
            }
        )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            [
                {
                    "fileName": item["fileName"],
                    "reportType": item["reportType"],
                    "paragraphCount": item["paragraphCount"],
                    "tableCount": item["tableCount"],
                }
                for item in inventory
            ],
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
