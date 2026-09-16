"""s0_register 集成测试(连真 PG;不可达 skip)。临时构造批次,测完按 FK 序清理。"""

import io
from datetime import date
from pathlib import Path

import pytest
from docx import Document as Docx
from openpyxl import Workbook
from sqlalchemy import delete, select, text
from ulid import ULID

from common.manifest import REQUIRED_COLUMNS
from common.pg_models import (
    Document,
    DocVersion,
    ImportBatch,
    PipelineEvent,
    RemediationRecord,
    ReviewQueue,
)
from pipeline.config import load_config
from pipeline.index.object_store import ObjectStore
from pipeline.index.pg_io import PgIO
from pipeline.queue import dispose
from pipeline.stage_base import StageContext
from pipeline.stages.s0_register import register_batch


@pytest.fixture
def pg():
    io_ = PgIO.from_config(load_config())
    try:
        with io_.session() as s:
            s.execute(text("select 1"))
    except Exception:
        pytest.skip("PG 不可达(demo up 未起)")
    return io_


@pytest.fixture
def reg(pg, tmp_path):
    ctx = StageContext(config=load_config(), object_store=ObjectStore(tmp_path / "obj"), db=pg)
    batches: list[str] = []
    yield ctx, tmp_path, batches
    with pg.session() as s:
        dvs = list(s.scalars(select(DocVersion).where(DocVersion.batch_id.in_(batches or [""]))))
        dvids = [dv.doc_version_id for dv in dvs]
        lids = {dv.logical_id for dv in dvs}
        if dvids:
            s.execute(delete(RemediationRecord).where(RemediationRecord.doc_version_id.in_(dvids)))
            s.execute(delete(ReviewQueue).where(ReviewQueue.doc_version_id.in_(dvids)))
            s.execute(delete(PipelineEvent).where(PipelineEvent.doc_version_id.in_(dvids)))
            s.execute(delete(DocVersion).where(DocVersion.doc_version_id.in_(dvids)))
        if lids:
            s.execute(delete(Document).where(Document.logical_id.in_(lids)))
        if batches:
            s.execute(delete(ImportBatch).where(ImportBatch.batch_id.in_(batches)))


def _docx(text_: str = "第一条 测试内容") -> bytes:
    buf = io.BytesIO()
    d = Docx()
    d.add_paragraph(text_)
    d.save(buf)
    return buf.getvalue()


def _pdf() -> bytes:
    return b"%PDF-1.4\n% minimal demo pdf\n%%EOF\n"


def _xlsx() -> bytes:
    buf = io.BytesIO()
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(buf)
    return buf.getvalue()


def _bid() -> str:
    return "t_" + str(ULID())  # 完整 ULID:前缀是时间戳,截断会在同毫秒内相撞


def _make_batch(
    base: Path, bid: str, rows: list[dict], files: dict, *, drop_col=None, extra_col=None
):
    d = base / bid
    d.mkdir(parents=True, exist_ok=True)
    for fn, data in files.items():
        (d / fn).write_bytes(data)
    cols = [c for c in REQUIRED_COLUMNS if c != drop_col]
    if extra_col:
        cols.append(extra_col)
    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    mp = d / "manifest.xlsx"
    wb.save(str(mp))
    return d, mp


def _row(filename, **kw):
    base = {
        "filename": filename, "title": f"标题-{filename}", "doc_number": f"令-{filename}",
        "issuer": "CSRC", "perm_tag": "公开", "corpus_type": "P-EXT",
        "biz_domain": "DISCLOSURE", "issue_date": "2024-01-01", "supersedes": "",
        "sub_type": "部门规章", "effective_date": "2024-01-01",
    }
    base.update(kw)
    return base


def test_manifest_missing_column_rejects_batch(reg):
    ctx, base, _ = reg
    bid = _bid()
    d, mp = _make_batch(base, bid, [_row("a.docx")], {"a.docx": _docx()}, drop_col="biz_domain")
    rep = register_batch(ctx, bid, d, mp)
    assert rep.accepted is False and "biz_domain" in rep.reject_reason
    assert ctx.db.get(ImportBatch, bid) is None  # 整批拒收,未落库


def test_manifest_extra_column_rejects_batch(reg):
    ctx, base, _ = reg
    bid = _bid()
    d, mp = _make_batch(
        base, bid, [_row("a.docx")], {"a.docx": _docx()}, extra_col="unexpected_col"
    )
    rep = register_batch(ctx, bid, d, mp)
    assert rep.accepted is False and "unexpected_col" in rep.reject_reason  # 多列整批拒收
    assert ctx.db.get(ImportBatch, bid) is None  # 未落库


def test_issue_date_written_to_doc_version(reg):
    ctx, base, batches = reg
    bid = _bid()
    batches.append(bid)
    d, mp = _make_batch(base, bid, [_row("a.docx", issue_date="2024-03-15")], {"a.docx": _docx()})
    o = register_batch(ctx, bid, d, mp).outcomes[0]
    dv = ctx.db.get(DocVersion, o.doc_version_id)
    assert dv.issue_date == date(2024, 3, 15)  # manifest issue_date 写入并归一为 date


def test_registers_docx_and_pdf(reg):
    ctx, base, batches = reg
    bid = _bid()
    batches.append(bid)
    rows = [_row("a.docx", corpus_type="P-INT", perm_tag="内部"), _row("b.pdf")]
    d, mp = _make_batch(base, bid, rows, {"a.docx": _docx(), "b.pdf": _pdf()})
    rep = register_batch(ctx, bid, d, mp)
    assert rep.accepted and rep.counts()["REGISTERED"] == 2
    for o in rep.outcomes:
        dv = ctx.db.get(DocVersion, o.doc_version_id)
        assert dv.pipeline_status == "REGISTERED"
        assert len(o.doc_version_id) == 26  # ULID
        assert ctx.object_store.exists(dv.raw_object_key)  # 原件写一次


def test_whitelist_quarantine(reg):
    # 白名单外格式(unknown 魔数)→ 隔离;xlsx 端到端留 P2 不入白名单,故用 unknown 测
    ctx, base, batches = reg
    bid = _bid()
    batches.append(bid)
    d, mp = _make_batch(
        base, bid, [_row("c.bin", corpus_type="P-INT")], {"c.bin": b"\x00\x01\x02 not whitelisted"}
    )
    o = register_batch(ctx, bid, d, mp).outcomes[0]
    assert o.status == "QUARANTINED" and o.error_code == "E101-DEMO"
    # B2:隔离件进统一队列(quarantine 类型),queue list 可见
    q = next(r for r in _queue(ctx, o.doc_version_id) if r.queue_type == "quarantine")
    assert q.status == "open" and q.evidence["error_code"] == "E101-DEMO"


def test_quarantine_release_reenters_parsing(reg):
    """B2 端到端:隔离件经统一队列 release → QUARANTINED 重入 PARSING。"""
    ctx, base, batches = reg
    bid = _bid()
    batches.append(bid)
    d, mp = _make_batch(base, bid, [_row("a.docx", perm_tag="")], {"a.docx": _docx()})
    o = register_batch(ctx, bid, d, mp).outcomes[0]
    q = next(r for r in _queue(ctx, o.doc_version_id) if r.queue_type == "quarantine")
    out = dispose(ctx.db, q.queue_id, "release", operator="test")
    assert out.after_state == "PARSING"
    assert ctx.db.get(DocVersion, o.doc_version_id).pipeline_status == "PARSING"
    closed = next(r for r in _queue(ctx, o.doc_version_id) if r.queue_id == q.queue_id)
    assert closed.status == "closed"


def test_perm_missing_quarantine(reg):
    ctx, base, batches = reg
    bid = _bid()
    batches.append(bid)
    d, mp = _make_batch(base, bid, [_row("a.docx", perm_tag="")], {"a.docx": _docx()})
    o = register_batch(ctx, bid, d, mp).outcomes[0]
    assert o.status == "QUARANTINED" and "密级缺失" in o.reason


def test_sha_dedup_second_is_duplicate(reg):
    ctx, base, batches = reg
    data = _docx("同一内容")
    bid1 = _bid()
    batches.append(bid1)
    d1, mp1 = _make_batch(base, bid1, [_row("a.docx")], {"a.docx": data})
    register_batch(ctx, bid1, d1, mp1)
    bid2 = _bid()
    batches.append(bid2)
    d2, mp2 = _make_batch(base, bid2, [_row("a.docx")], {"a.docx": data})
    o = register_batch(ctx, bid2, d2, mp2).outcomes[0]
    assert o.status == "DUPLICATE"
    # 重复登记在既有 doc 留审计事件(report 未持久化,否则去重关联 DB 无痕)
    with ctx.db.session() as s:
        evs = list(
            s.scalars(select(PipelineEvent).where(PipelineEvent.doc_version_id == o.doc_version_id))
        )
    dup = [e for e in evs if e.detail and "duplicate_ingest" in e.detail]
    assert len(dup) == 1
    assert dup[0].detail["duplicate_ingest"] == {"batch_id": bid2, "filename": "a.docx"}
    assert dup[0].from_state == dup[0].to_state  # 非迁移审计记录


def test_reingest_same_batch_id_is_idempotent(reg):
    ctx, base, batches = reg
    bid = _bid()
    batches.append(bid)
    d, mp = _make_batch(base, bid, [_row("a.docx")], {"a.docx": _docx("同一内容")})
    out1 = register_batch(ctx, bid, d, mp).outcomes[0]
    assert out1.status == "REGISTERED"
    # 同 batch_id 重跑:不撞主键(get-or-create),SHA 去重命中 → DUPLICATE,doc_version_id 不变
    out2 = register_batch(ctx, bid, d, mp).outcomes[0]
    assert out2.status == "DUPLICATE"
    assert out2.doc_version_id == out1.doc_version_id  # 复用既有 doc_version → chunk_id 稳定
    with ctx.db.session() as s:  # 仍只有一行批次
        n = len(list(s.scalars(select(ImportBatch).where(ImportBatch.batch_id == bid))))
    assert n == 1


def test_supersede_inherits_logical(reg):
    ctx, base, batches = reg
    bid1 = _bid()
    batches.append(bid1)
    d1, mp1 = _make_batch(base, bid1, [_row("v1.docx", doc_number="令1")], {"v1.docx": _docx("v1")})
    out1 = register_batch(ctx, bid1, d1, mp1).outcomes[0]
    bid2 = _bid()
    batches.append(bid2)
    r2 = [_row("v2.docx", doc_number="令2", supersedes="v1.docx")]
    d2, mp2 = _make_batch(base, bid2, r2, {"v2.docx": _docx("v2")})
    out2 = register_batch(ctx, bid2, d2, mp2).outcomes[0]
    assert out2.logical_id == out1.logical_id  # 替代 → logical 继承
    dv2 = ctx.db.get(DocVersion, out2.doc_version_id)
    assert dv2.supersedes_version_id == out1.doc_version_id
    assert dv2.version_relation == "revise_replace"


def _queue(ctx, dvid):
    with ctx.db.session() as s:
        return list(s.scalars(select(ReviewQueue).where(ReviewQueue.doc_version_id == dvid)))


def test_supersede_abolish_only_new_logical(reg):
    ctx, base, batches = reg
    bid1 = _bid()
    batches.append(bid1)
    d1, mp1 = _make_batch(base, bid1, [_row("v1.docx", doc_number="令1")], {"v1.docx": _docx("v1")})
    out1 = register_batch(ctx, bid1, d1, mp1).outcomes[0]
    bid2 = _bid()
    batches.append(bid2)
    r2 = [_row("notice.docx", doc_number="令2", supersedes="abolish:v1.docx")]
    d2, mp2 = _make_batch(base, bid2, r2, {"notice.docx": _docx("废止公告")})
    out2 = register_batch(ctx, bid2, d2, mp2).outcomes[0]
    assert out2.logical_id != out1.logical_id  # abolish_only 不继承 logical(独立文书)
    dv2 = ctx.db.get(DocVersion, out2.doc_version_id)
    assert dv2.version_relation == "abolish_only"
    assert dv2.supersedes_version_id == out1.doc_version_id  # 记被废止版


def test_supersede_merge_enqueues_meta_confirm(reg):
    ctx, base, batches = reg
    bid1 = _bid()
    batches.append(bid1)
    rows1 = [_row("a.docx", doc_number="令A"), _row("b.docx", doc_number="令B")]
    d1, mp1 = _make_batch(base, bid1, rows1, {"a.docx": _docx("aaa"), "b.docx": _docx("bbb")})
    register_batch(ctx, bid1, d1, mp1)
    bid2 = _bid()
    batches.append(bid2)
    r2 = [_row("merged.docx", doc_number="令M", supersedes="a.docx;b.docx")]
    d2, mp2 = _make_batch(base, bid2, r2, {"merged.docx": _docx("merged")})
    out2 = register_batch(ctx, bid2, d2, mp2).outcomes[0]
    assert ctx.db.get(DocVersion, out2.doc_version_id).version_relation is None  # 不自动建模
    qs = _queue(ctx, out2.doc_version_id)
    assert len(qs) == 1 and qs[0].queue_type == "meta_confirm"
    assert qs[0].evidence["relation"] == "merge"


def test_supersede_split_enqueues_meta_confirm(reg):
    ctx, base, batches = reg
    bid1 = _bid()
    batches.append(bid1)
    d1, mp1 = _make_batch(
        base, bid1, [_row("orig.docx", doc_number="令O")], {"orig.docx": _docx("orig")}
    )
    register_batch(ctx, bid1, d1, mp1)
    bid2 = _bid()  # 同批两新件都 supersede orig → split
    batches.append(bid2)
    r2 = [
        _row("p1.docx", doc_number="令P1", supersedes="orig.docx"),
        _row("p2.docx", doc_number="令P2", supersedes="orig.docx"),
    ]
    d2, mp2 = _make_batch(base, bid2, r2, {"p1.docx": _docx("p1"), "p2.docx": _docx("p2")})
    outs = register_batch(ctx, bid2, d2, mp2).outcomes
    for o in outs:
        assert ctx.db.get(DocVersion, o.doc_version_id).version_relation is None
        qs = _queue(ctx, o.doc_version_id)
        assert len(qs) == 1 and qs[0].evidence["relation"] == "split_replace"
