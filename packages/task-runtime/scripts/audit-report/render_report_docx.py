from __future__ import annotations

import argparse
import json
import os
import re
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from lxml import etree

WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def first_text_run(paragraph: Any) -> Any:
    run = next((item for item in paragraph.runs if item.text.strip()), None)
    if run is None:
        raise ValueError(f"Template prototype has no formatted run: {paragraph.text!r}")
    return run


def clone_paragraph(document: DocumentObject, prototype: Any, text: str) -> Any:
    if prototype is None:
        raise ValueError(f"Template does not provide a prototype for: {text[:40]!r}")
    paragraph = document.add_paragraph()
    try:
        paragraph.style = prototype.style
    except (KeyError, ValueError):
        pass
    if prototype._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(prototype._p.pPr))
    run = paragraph.add_run(text)
    source_run = first_text_run(prototype)
    if source_run._r.rPr is not None:
        run._r.insert(0, deepcopy(source_run._r.rPr))
    return paragraph


def find_prototypes(document: DocumentObject) -> dict[str, Any]:
    paragraphs = [item for item in document.paragraphs if item.text.strip()]
    if len(paragraphs) < 4:
        raise ValueError("Template must contain formatted title, body, closing and date prototypes")

    def find(pattern: str) -> Any | None:
        return next((item for item in paragraphs if re.search(pattern, item.text.strip())), None)

    date = find(r"^\d{4}年.+日$")
    if date is None:
        raise ValueError("Template does not contain a formatted report-date prototype")
    date_index = paragraphs.index(date)
    closing = paragraphs[date_index - 1] if date_index > 0 else None
    body = find(r"^(?:按照审计工作安排|根据.+委托)")
    prototypes = {
        "title-1": paragraphs[0],
        "title-2": paragraphs[1] if len(paragraphs) > 1 else paragraphs[0],
        "addressee": find(r"[：:]$"),
        "body": body,
        "section": find(r"^[一二三四五六七八九十]+、"),
        "subsection": find(r"^（[一二三四五六七八九十]+）"),
        "finding-title": find(r"^\d+[.．、]"),
        "note": find(r"^备注："),
        "table-title": find(r"^表\d+"),
        "closing": closing,
        "date": date,
    }
    for key in ("body", "section", "closing", "date"):
        if prototypes[key] is None:
            raise ValueError(f"Template does not contain required prototype: {key}")
    return prototypes


def clear_body(document: DocumentObject) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clone_cell_format(target: Any, source: Any, value: Any) -> None:
    target.text = ""
    if source._tc.tcPr is not None:
        target._tc.insert(0, deepcopy(source._tc.tcPr))
    source_paragraph = source.paragraphs[0]
    target_paragraph = target.paragraphs[0]
    if source_paragraph._p.pPr is not None:
        target_paragraph._p.insert(0, deepcopy(source_paragraph._p.pPr))
    run = target_paragraph.add_run(str(value))
    # Some official table prototypes intentionally keep data cells empty. In
    # that case paragraph/cell properties still carry the layout, while there
    # is no non-empty run whose character properties can be copied.
    source_run = next((item for item in source_paragraph.runs if item.text.strip()), None)
    if source_run is None:
        source_run = next(iter(source_paragraph.runs), None)
    if source_run is not None and source_run._r.rPr is not None:
        run._r.insert(0, deepcopy(source_run._r.rPr))
    if isinstance(value, (int, float)) and value < 0:
        run.font.color.rgb = RGBColor(255, 0, 0)


def add_table(document: DocumentObject, data: dict[str, Any], title_prototype: Any, prototype: Any) -> None:
    if prototype is None:
        raise ValueError(f"Template has no table prototype for {data['tableId']}")
    title = data["title"]
    if data.get("unit"):
        title = f"{title}\t单位：{data['unit']}"
    clone_paragraph(document, title_prototype, title)
    headers = data["headers"]
    table = document.add_table(rows=1, cols=len(headers))
    if prototype._tbl.tblPr is not None:
        table._tbl.remove(table._tbl.tblPr)
        table._tbl.insert(0, deepcopy(prototype._tbl.tblPr))
    # Reuse the official column grid only when the generated table has the
    # same number of columns. Copying a five-column prototype grid into a
    # three-column period table leaves visible blank columns in Word.
    if prototype._tbl.tblGrid is not None and len(prototype.columns) == len(headers):
        table._tbl.remove(table._tbl.tblGrid)
        table._tbl.insert(1, deepcopy(prototype._tbl.tblGrid))
    for index, header in enumerate(headers):
        clone_cell_format(table.rows[0].cells[index], prototype.rows[0].cells[min(index, len(prototype.columns) - 1)], header)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    source_row = prototype.rows[min(1, len(prototype.rows) - 1)]
    for values in data["rows"]:
        row = table.add_row()
        for index, value in enumerate(values):
            clone_cell_format(row.cells[index], source_row.cells[min(index, len(prototype.columns) - 1)], value)


def add_paragraph(document: DocumentObject, item: dict[str, Any], prototypes: dict[str, Any]) -> None:
    if item["paragraphId"].endswith("-title"):
        key = "finding-title"
    elif item["text"].startswith("备注："):
        key = "note"
    else:
        key = "body"
    # A report template may not contain an example finding title or note when
    # that particular sample has no findings. Preserve report generation by
    # falling back to the template's body style for these optional prototypes.
    clone_paragraph(document, prototypes.get(key) or prototypes["body"], item["text"])


def add_draft(document: DocumentObject, draft: dict[str, Any], prototypes: dict[str, Any], table_prototype: Any) -> None:
    for index, title in enumerate(draft["titleLines"]):
        clone_paragraph(document, prototypes[f"title-{min(index + 1, 2)}"], title)
    if draft.get("addressee"):
        clone_paragraph(document, prototypes["addressee"], draft["addressee"])
    clone_paragraph(document, prototypes["body"], draft["introduction"]["text"])
    for section in draft["sections"]:
        clone_paragraph(document, prototypes["section"], section["heading"])
        for item in section["paragraphs"]:
            add_paragraph(document, item, prototypes)
        for table in section["tables"]:
            add_table(document, table, prototypes["table-title"], table_prototype)
        for subsection in section["subsections"]:
            clone_paragraph(document, prototypes["subsection"], subsection["heading"])
            split = subsection.get("tablesAfterParagraphCount", len(subsection["paragraphs"]))
            for index, item in enumerate(subsection["paragraphs"]):
                if index == split:
                    for table in subsection.get("tables", []):
                        add_table(document, table, prototypes["table-title"], table_prototype)
                add_paragraph(document, item, prototypes)
            if split >= len(subsection["paragraphs"]):
                for table in subsection.get("tables", []):
                    add_table(document, table, prototypes["table-title"], table_prototype)
        for item in section.get("closingParagraphs", []):
            add_paragraph(document, item, prototypes)
    clone_paragraph(document, prototypes["closing"], draft["closingOrganization"])
    clone_paragraph(document, prototypes["date"], draft["reportDate"])
    if draft["status"] == "needs-input":
        clone_paragraph(document, prototypes["body"], "【数据未就绪：本草稿不得作为正式审计报告】")


def sanitize_review_markup(path: Path) -> None:
    removable_parts = {"word/comments.xml", "word/commentsExtended.xml", "word/commentsExtensible.xml", "word/commentsIds.xml", "word/people.xml"}
    removable_tags = {f"{{{WORD_NAMESPACE}}}{name}" for name in ("commentRangeStart", "commentRangeEnd", "commentReference", "del", "moveFrom", "pPrChange", "rPrChange", "tblPrChange", "tblGridChange", "trPrChange", "tcPrChange", "sectPrChange", "numPrChange")}
    accepted_tags = {f"{{{WORD_NAMESPACE}}}{name}" for name in ("ins", "moveTo")}
    with tempfile.NamedTemporaryFile(suffix=".docx", dir=path.parent, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                if info.filename in removable_parts:
                    continue
                payload = source.read(info.filename)
                if info.filename.endswith((".xml", ".rels")):
                    root = etree.fromstring(payload)
                    changed = False
                    for element in list(root.iter()):
                        parent = element.getparent()
                        if parent is None:
                            continue
                        if element.tag in removable_tags:
                            parent.remove(element)
                            changed = True
                        elif element.tag in accepted_tags:
                            position = parent.index(element)
                            for child in list(element):
                                parent.insert(position, child)
                                position += 1
                            parent.remove(element)
                            changed = True
                        elif any(token in element.get("Target", "") or token in element.get("PartName", "") for token in ("comments", "people")):
                            parent.remove(element)
                            changed = True
                    if changed:
                        payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                target.writestr(info, payload)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Render ReportDraft JSON by inheriting an external DOCX template")
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--draft", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    draft = json.loads(args.draft.read_text(encoding="utf-8"))
    document = Document(str(args.template))
    prototypes = find_prototypes(document)
    table_prototype = document.tables[0] if document.tables else None
    clear_body(document)
    add_draft(document, draft, prototypes, table_prototype)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(args.output))
    sanitize_review_markup(args.output)


if __name__ == "__main__":
    main()
