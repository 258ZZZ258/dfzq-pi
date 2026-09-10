from __future__ import annotations

import importlib.util
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SCRIPT = Path(__file__).with_name("render_supervision_narrative_docx.py")
SPEC = importlib.util.spec_from_file_location("render_supervision_narrative_docx", SCRIPT)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class RenderSupervisionNarrativeDocxTest(unittest.TestCase):
    def test_renders_narrative_report_without_prose_tables(self) -> None:
        analysis = {
            "schemaVersion": "supervision-analysis.v1",
            "task": {
                "analysisPeriod": {"from": "2024-01-01", "to": "2024-12-31"},
                "organizationId": "东方证券股份有限公司江阴人民东路证券营业部",
            },
            "snapshot": {
                "included": [
                    {
                        "documentId": "DOC-1",
                        "title": "监管决定书",
                        "sourceType": "regulatory",
                        "fileDate": "2024-01-01",
                        "organizationIds": ["东方证券股份有限公司"],
                        "dataOrigin": "official-public",
                        "uploadEntry": "supervision",
                        "sourceUrl": "https://example.com/source",
                    }
                ]
            },
            "issues": [
                {
                    "issueId": "ISSUE-1",
                    "sourceDocumentId": "DOC-1",
                    "dataOrigin": "official-public",
                    "category": "业务管理",
                    "description": "测试问题描述。",
                    "evidenceIds": ["DOC-1:p1:issue-1"],
                }
            ],
            "relations": [
                {
                    "issueId": "ISSUE-1",
                    "relationType": "RECTIFICATION",
                    "status": "UNMATCHED",
                    "selectedRecordId": None,
                }
            ],
        }
        narrative = {
            "schemaVersion": "supervision-report-narrative.v2",
            "executiveSummary": "报告期内，营业部接受外部监管检查，相关监管文书反映业务管理环节存在控制执行不到位的情形。营业部围绕问题组织整改，进一步完善制度执行、过程复核和资料留存要求。内部监督持续关注重点业务环节，督促责任岗位落实管理要求。",
            "regulatoryOverview": "报告期内，监管机构对营业部开展现场检查并就检查发现事项采取监管措施。有关事项主要涉及业务流程执行和日常管理控制，反映出部分制度要求在具体经营活动中落实不够充分。营业部已围绕监管指出问题组织专项整改。",
            "regulatoryIssues": [{"title": "业务管理事项", "analysis": "监管检查发现，营业部在相关业务管理过程中存在操作程序执行不够规范的情况，反映出前置审核、过程复核和档案留存等控制要求尚未完全落实。该问题不利于业务活动全过程的规范管理和责任追溯，应当通过完善操作标准、强化复核监督等方式予以纠正。", "issueIds": ["ISSUE-1"]}],
            "regulatoryRectification": "针对监管指出的业务管理问题，营业部组织开展流程梳理和制度核对，对相关岗位提出整改要求，明确业务办理、复核和资料归档的操作标准。同时加强日常检查和执行监督，推动整改措施落实到具体业务环节，促进相关管理要求形成常态化机制。",
            "externalAuditAnalysis": "报告期内，外部检查重点关注业务管理流程和资料留存情况。检查发现的事项表明，营业部仍需进一步强化关键环节的操作规范，加强重要业务记录的完整性管理，并通过持续复核确保相关制度要求得到有效执行。",
            "internalInspectionOverview": "报告期内，内部审计、合规和风险检查围绕制度执行、业务操作和风险控制开展监督，对经营管理中的薄弱环节进行梳理，并督促有关岗位落实整改要求。内部监督发现事项总体与外部监管关注方向保持一致。",
            "internalInspectionThemes": [{"title": "制度及内控机制建设", "analysis": "内部检查发现，个别业务环节的制度执行和过程留痕仍不够规范，反映出控制要求向具体岗位和操作动作转化不够充分。营业部已结合检查意见完善相关操作要求，进一步明确审核、复核和归档责任，并将执行情况纳入日常监督。", "issueIds": ["ISSUE-1"]}],
            "accountabilityAnalysis": "报告期内，营业部依据内部管理要求对有关事项开展责任核查，并对相关责任岗位进行合规提醒，督促其强化履职意识、严格落实业务流程和合规审查要求。问责处理与问题整改同步推进，发挥了警示和教育作用。",
            "violationAccountabilityAnalysis": "报告期内未发生需要启动违规经营投资责任追究程序的事项。",
            "routineComplianceAnalysis": "合规管理部门围绕重点业务事项持续开展提示和督导，要求有关岗位严格执行前置审查、客户适当性管理和业务资料留存要求，并按规定反馈自查整改情况，推动合规要求嵌入日常经营流程。",
            "routineRiskAnalysis": "风险管理部门持续关注营业终端、人员使用和监控接入情况，督促营业部加强资产台账与监控登记信息核对，确保终端管理责任清晰、使用状态可查、操作过程可追溯。",
            "litigationAnalysis": "报告期内涉及一宗证券交易委托纠纷，案件处于一审审理阶段。营业部已将该事项纳入法律事务跟踪范围，持续关注案件进展，并结合争议事项检查客户沟通、委托指令和业务留痕管理情况。",
            "generation": {"provider": "deepseek", "model": "deepseek-v4-flash"},
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "report.docx"
            document = MODULE.build_document(analysis, narrative, {}, None, "2026年9月2日")
            document.save(output)
            rendered = Document(output)
            text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
            self.assertIn("监督信息汇总分析报告", text)
            self.assertIn("一、摘要", text)
            self.assertIn("二、外部事项", text)
            self.assertIn("（一）监管检查、函件等指出问题及整改情况", text)
            self.assertIn("（二）外部审计、检查、调查等指出问题及整改情况", text)
            self.assertIn("三、内部事项", text)
            self.assertIn("（一）内部审计、检查发现的问题及整改情况（审计、合规、风险）", text)
            self.assertIn("（二）问责处理情况", text)
            self.assertIn("（三）日常监督情况", text)
            self.assertIn("附件", text)
            self.assertIn("3.法律诉讼案件", text)
            self.assertIn(
                "被分析单位：东方证券股份有限公司江阴人民东路证券营业部；简称：江阴人民东路营业部",
                text,
            )
            self.assertEqual(text.count("被分析单位："), 1)
            self.assertIn("监管检查发现，营业部在相关业务管理过程中", text)
            self.assertNotIn("尚缺证据", text)
            self.assertNotIn("待业务确认", text)
            self.assertNotIn("本批次", text)
            paragraphs = [paragraph for paragraph in rendered.paragraphs if paragraph.text.strip()]
            framework_index = next(
                index for index, paragraph in enumerate(paragraphs) if paragraph.text == "1.业务管理事项"
            )
            self.assertEqual(paragraphs[framework_index].style.name, "Framework Heading")
            self.assertAlmostEqual(
                paragraphs[framework_index].style.paragraph_format.left_indent.cm,
                0.74,
                places=2,
            )
            self.assertTrue(paragraphs[framework_index + 1].text.startswith("监管检查发现"))
            level_two = next(
                paragraph
                for paragraph in paragraphs
                if paragraph.text == "（一）监管检查、函件等指出问题及整改情况"
            )
            self.assertEqual(level_two.style.name, "Heading 2")
            chinese_run = next(run for run in level_two.runs if "监管检查" in run.text)
            punctuation_run = level_two.runs[0]
            self.assertEqual(chinese_run._element.rPr.rFonts.get(qn("w:eastAsia")), "黑体")
            self.assertEqual(punctuation_run._element.rPr.rFonts.get(qn("w:eastAsia")), "Times New Roman")
            self.assertEqual(punctuation_run._element.rPr.rFonts.get(qn("w:ascii")), "Times New Roman")
            self.assertNotIn("资料来源", text)
            self.assertEqual(len(rendered.tables), 0)
            hyperlinks = [rel for rel in rendered.part.rels.values() if rel.reltype.endswith("/hyperlink")]
            self.assertEqual(hyperlinks, [])
            with zipfile.ZipFile(output) as package:
                visible_xml = package.read("word/document.xml") + package.read("word/styles.xml")
            self.assertNotIn(b"0563C1", visible_xml)
            self.assertAlmostEqual(rendered.sections[0].page_width.cm, 21.0, places=1)


if __name__ == "__main__":
    unittest.main()
