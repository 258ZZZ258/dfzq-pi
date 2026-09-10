from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(100, 100, 100)

# `standard_business_brief` with named Chinese-report overrides:
# A4 paper, 3 cm side margins, Chinese official-document fonts, black headings,
# restrained spacing, and prose-only body sections.


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def set_run_font(
    run: Any,
    *,
    name: str,
    size: float,
    bold: bool = False,
    color: RGBColor = BLACK,
    latin_name: str = "Times New Roman",
) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin_name)


def set_spacing(paragraph: Any, *, before: float = 0, after: float = 6, line: float = 1.5) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, name="宋体", size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, name="宋体", size=9, color=MUTED)


def configure_document() -> DocumentObject:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.2)
    add_page_number(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.45

    heading_specs = {
        "Heading 1": ("黑体", 15, 12, 6),
        "Heading 2": ("黑体", 13, 9, 4),
        "Heading 3": ("黑体", 11, 6, 3),
    }
    for style_name, (font_name, font_size, before, after) in heading_specs.items():
        style = document.styles[style_name]
        style.font.name = font_name
        style.font.size = Pt(font_size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        heading_fonts = style._element.get_or_add_rPr().rFonts
        heading_fonts.set(qn("w:eastAsia"), font_name)
        heading_fonts.set(qn("w:ascii"), "Times New Roman")
        heading_fonts.set(qn("w:hAnsi"), "Times New Roman")
        heading_fonts.set(qn("w:cs"), "Times New Roman")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.keep_with_next = True
        if style_name == "Heading 3":
            style.paragraph_format.left_indent = Cm(0.74)

    framework_heading = document.styles.add_style("Framework Heading", WD_STYLE_TYPE.PARAGRAPH)
    framework_heading.font.name = "黑体"
    framework_heading.font.size = Pt(10.5)
    framework_heading.font.bold = True
    framework_heading.font.color.rgb = BLACK
    framework_fonts = framework_heading._element.get_or_add_rPr().rFonts
    framework_fonts.set(qn("w:eastAsia"), "黑体")
    framework_fonts.set(qn("w:ascii"), "Times New Roman")
    framework_fonts.set(qn("w:hAnsi"), "Times New Roman")
    framework_fonts.set(qn("w:cs"), "Times New Roman")
    framework_heading.paragraph_format.left_indent = Cm(0.74)
    framework_heading.paragraph_format.first_line_indent = Cm(0)
    framework_heading.paragraph_format.space_before = Pt(4)
    framework_heading.paragraph_format.space_after = Pt(2)
    framework_heading.paragraph_format.line_spacing = 1.2
    framework_heading.paragraph_format.keep_with_next = True

    note = document.styles.add_style("Report Note", WD_STYLE_TYPE.PARAGRAPH)
    note.font.name = "宋体"
    note.font.size = Pt(9.5)
    note.font.color.rgb = BLACK
    note._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    note.paragraph_format.left_indent = Cm(0.74)
    note.paragraph_format.right_indent = Cm(0)
    note.paragraph_format.space_before = Pt(1)
    note.paragraph_format.space_after = Pt(5)
    note.paragraph_format.line_spacing = 1.25

    item = document.styles.add_style("Report Item", WD_STYLE_TYPE.PARAGRAPH)
    item.font.name = "宋体"
    item.font.size = Pt(10.5)
    item._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    item.paragraph_format.left_indent = Cm(0.74)
    item.paragraph_format.first_line_indent = Cm(-0.74)
    item.paragraph_format.space_before = Pt(2)
    item.paragraph_format.space_after = Pt(3)
    item.paragraph_format.line_spacing = 1.35
    return document


def add_heading(document: DocumentObject, text: str, level: int) -> Any:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    add_mixed_heading_runs(
        paragraph,
        text,
        chinese_font="黑体",
        size={1: 15, 2: 13, 3: 11}[level],
    )
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_mixed_heading_runs(
    paragraph: Any,
    text: str,
    *,
    chinese_font: str,
    size: float,
) -> None:
    """Use Chinese font for Han characters and Times New Roman for all other title glyphs."""
    for part in re.findall(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]+|[^\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]+", text):
        is_chinese = bool(re.fullmatch(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]+", part))
        font_name = chinese_font if is_chinese else "Times New Roman"
        run = paragraph.add_run(part)
        set_run_font(run, name=font_name, size=size, bold=True, latin_name="Times New Roman")


def add_body(document: DocumentObject, text: str, *, indent: bool = True) -> Any:
    paragraph = document.add_paragraph()
    if not indent:
        paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(text)
    set_run_font(run, name="宋体", size=11)
    return paragraph


def add_note(document: DocumentObject, text: str) -> Any:
    paragraph = document.add_paragraph(style="Report Note")
    run = paragraph.add_run(text)
    set_run_font(run, name="宋体", size=9.5, color=BLACK)
    return paragraph


def add_item(document: DocumentObject, label: str, text: str, *, evidence: str | None = None) -> Any:
    paragraph = document.add_paragraph(style="Report Item")
    label_run = paragraph.add_run(f"{label}：")
    set_run_font(label_run, name="黑体", size=10.5, bold=True)
    value_run = paragraph.add_run(text)
    set_run_font(value_run, name="宋体", size=10.5)
    if evidence:
        evidence_run = paragraph.add_run(f"（证据：{evidence}）")
        set_run_font(evidence_run, name="宋体", size=9, color=MUTED)
    return paragraph


def add_title_block(document: DocumentObject, analysis: dict[str, Any], narrative: dict[str, Any], report_date: str) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("监督信息汇总分析报告")
    set_run_font(run, name="黑体", size=20, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(4)
    run = subtitle.add_run(analysis["task"]["organizationId"])
    set_run_font(run, name="宋体", size=12, color=BLACK)

    task = analysis["task"]
    period = task.get("analysisPeriod") or {"from": task["analysisStart"], "to": task["analysisEnd"]}
    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(metadata, after=4, line=1.2)
    run = metadata.add_run(
        f"分析期间：{period['from']} 至 {period['to']}    生成日期：{report_date}"
    )
    set_run_font(run, name="宋体", size=9.5, color=BLACK)

    property_line = document.add_paragraph()
    property_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(property_line, after=10, line=1.2)
    run = property_line.add_run("报告属性：业务验收测试稿（含模拟内部资料）")
    set_run_font(run, name="宋体", size=9, color=MUTED)


def relation_maps(analysis: dict[str, Any], records: dict[str, Any]) -> tuple[dict[tuple[str, str], dict[str, Any]], dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    relations = {(item["issueId"], item["relationType"]): item for item in analysis["relations"]}
    rectifications = {item["recordId"]: item for item in records.get("rectifications", [])}
    accountabilities = {item["recordId"]: item for item in records.get("accountabilities", [])}
    return relations, rectifications, accountabilities


def status_zh(value: str | None) -> str:
    return {
        "COMPLETED": "已完成",
        "IN_PROGRESS": "进行中",
        "PARTIALLY_COMPLETED": "部分完成",
        "ONGOING": "持续整改",
        "CONTINUOUS": "持续整改",
        "AUTO_CONFIRMED": "自动确认",
        "UNMATCHED": "待业务确认",
        "REVIEW_REQUIRED": "待复核",
    }.get(value or "", value or "未提供")


def first_sentences(text: str, count: int) -> str:
    sentences = [part.strip() for part in re.findall(r"[^。！？]+[。！？]", text) if part.strip()]
    if not sentences:
        return text.strip()
    return "".join(sentences[:count])


def date_zh(value: str) -> str:
    match = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", value)
    if not match:
        return value
    year, month, day = match.groups()
    return f"{year}年{int(month)}月{int(day)}日"


def sentence_fragment(value: Any) -> str:
    return str(value).strip().rstrip("。；; ")


def add_framework_item(document: DocumentObject, label: str, text: str) -> Any:
    """Render a framework title and its narrative as two distinct paragraphs."""
    title = document.add_paragraph(style="Framework Heading")
    add_mixed_heading_runs(title, label, chinese_font="黑体", size=10.5)
    body = add_body(document, text)
    return body


def add_scope_line(document: DocumentObject, text: str) -> Any:
    paragraph = add_body(document, text, indent=False)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def source_note_for_issues(
    document: DocumentObject,
    issues: list[dict[str, Any]],
    material_by_id: dict[str, dict[str, Any]],
) -> None:
    if not issues:
        return
    citations: list[str] = []
    seen: set[str] = set()
    for issue in issues:
        material = material_by_id[issue["sourceDocumentId"]]
        document_id = material["documentId"]
        if document_id in seen:
            continue
        seen.add(document_id)
        publisher = material.get("publisher") or "公开监管文书"
        document_number = issue.get("documentNumber") or "文号未载明"
        evidence = issue.get("fieldValues", {}).get("evidenceLocation") or "页码未载明"
        citations.append(
            f"{publisher}《{material['title']}》（{document_number}，"
            f"{material.get('fileDate', '日期未载明')}，{evidence}）"
        )
    add_note(document, "依据：" + "；".join(citations) + "。")


def unique_selected_records(
    issues: list[dict[str, Any]],
    relations: dict[tuple[str, str], dict[str, Any]],
    relation_type: str,
    records: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    seen: set[str] = set()
    for issue in issues:
        relation = relations.get((issue["issueId"], relation_type))
        record_id = relation.get("selectedRecordId") if relation else None
        if not record_id or record_id in seen or record_id not in records:
            continue
        seen.add(record_id)
        selected.append(records[record_id])
    return selected


def build_document(
    analysis: dict[str, Any],
    narrative: dict[str, Any],
    records: dict[str, Any],
    ocr: dict[str, Any] | None,
    report_date: str,
) -> DocumentObject:
    if analysis.get("schemaVersion") != "supervision-analysis.v1":
        raise ValueError("analysis schemaVersion must be supervision-analysis.v1")
    if narrative.get("schemaVersion") != "supervision-report-narrative.v2":
        raise ValueError("narrative schemaVersion must be supervision-report-narrative.v2")
    if narrative.get("generation", {}).get("provider") != "deepseek":
        raise ValueError("narrative must be generated by the project DeepSeek pipeline")

    document = configure_document()
    add_title_block(document, analysis, narrative, report_date)

    # The heading ladder below mirrors the workbook's “报告结构” sheet exactly.
    add_heading(document, "一、摘要", 1)
    add_body(document, narrative["executiveSummary"])

    add_heading(document, "二、外部事项", 1)
    add_heading(document, "（一）监管检查、函件等指出问题及整改情况", 2)
    add_scope_line(document, "（证监局、证券业协会、证券交易所等行业主管单位）")
    add_heading(document, "1.监管事项概况", 3)
    add_body(document, narrative["regulatoryOverview"])
    add_heading(document, "2.主要问题", 3)
    for index, theme in enumerate(narrative["regulatoryIssues"], start=1):
        add_framework_item(document, f"{index}.{theme['title']}", theme["analysis"])
    add_heading(document, "3.整改情况", 3)
    add_body(document, narrative["regulatoryRectification"])

    add_heading(document, "（二）外部审计、检查、调查等指出问题及整改情况", 2)
    add_scope_line(document, "（审计署、审计局、国资委、申能集团等上级单位）")
    add_body(document, narrative["externalAuditAnalysis"])

    add_heading(document, "三、内部事项", 1)
    add_heading(document, "（一）内部审计、检查发现的问题及整改情况（审计、合规、风险）", 2)
    add_body(document, narrative["internalInspectionOverview"])
    for index, theme in enumerate(narrative["internalInspectionThemes"], start=1):
        add_framework_item(document, f"{index}.{theme['title']}", theme["analysis"])

    add_heading(document, "（二）问责处理情况", 2)
    add_framework_item(document, "1.合规问责", narrative["accountabilityAnalysis"])
    add_framework_item(
        document,
        "2.违规经营投资责任追究",
        narrative["violationAccountabilityAnalysis"],
    )

    add_heading(document, "（三）日常监督情况", 2)
    add_framework_item(
        document,
        "1.合规警示/提示/建议/关注/问询",
        narrative["routineComplianceAnalysis"],
    )
    add_framework_item(document, "2.风险关注/提示/建议", narrative["routineRiskAnalysis"])
    add_framework_item(document, "3.法律诉讼案件", narrative["litigationAnalysis"])

    add_heading(document, "附件", 1)
    abbreviations = {
        "东方证券股份有限公司": "东方证券",
        "东方金融控股（香港）有限公司": "东方香港",
        "东方证券承销保荐有限公司": "东方证券承销保荐",
        "东方证券股份有限公司江阴人民东路证券营业部": "江阴人民东路营业部",
        "东方证券股份有限公司沈阳南八中路证券营业部": "沈阳南八中路营业部",
    }
    task = analysis["task"]
    organization = task["organizationId"]
    abbreviation = abbreviations.get(organization, organization)
    period = task.get("analysisPeriod") or {"from": task["analysisStart"], "to": task["analysisEnd"]}
    add_body(
        document,
        f"被分析单位：{organization}；简称：{abbreviation}。分析区间：{period['from']}至{period['to']}。",
        indent=False,
    )
    return document


def main() -> None:
    parser = argparse.ArgumentParser(description="Render a DeepSeek supervision narrative as a formal DOCX report")
    parser.add_argument("--analysis", required=True, type=Path)
    parser.add_argument("--narrative", required=True, type=Path)
    parser.add_argument("--records", required=True, type=Path)
    parser.add_argument("--ocr", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report-date")
    args = parser.parse_args()
    report_date = args.report_date or f"{date.today().year}年{date.today().month}月{date.today().day}日"
    document = build_document(
        read_json(args.analysis),
        read_json(args.narrative),
        read_json(args.records),
        read_json(args.ocr) if args.ocr else None,
        report_date,
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(json.dumps({"output": str(args.output), "narrativeProvider": "deepseek"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
