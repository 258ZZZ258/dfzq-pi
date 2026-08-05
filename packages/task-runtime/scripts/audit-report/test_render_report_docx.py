from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt


class RenderReportDocxTest(unittest.TestCase):
    def test_inherits_external_template_and_writes_draft(self) -> None:
        script = Path(__file__).with_name("render_report_docx.py")
        with tempfile.TemporaryDirectory() as raw_dir:
            root = Path(raw_dir)
            template_path = root / "template.docx"
            draft_path = root / "draft.json"
            output_path = root / "output.docx"

            template = Document()
            for text, size in (
                ("XX证券营业部", 22),
                ("常规审计报告", 22),
                ("XX证券营业部：", 16),
                ("按照审计工作安排，模板正文。", 16),
                ("一、审计项目基本情况", 16),
                ("（一）经营情况", 16),
                ("1.问题标题", 16),
                ("备注：模板说明", 10.5),
                ("表1 经营情况", 11),
                ("模板落款机构", 16),
                ("2026年1月1日", 16),
            ):
                paragraph = template.add_paragraph()
                run = paragraph.add_run(text)
                run.font.size = Pt(size)
            template.save(template_path)

            draft = {
                "taskId": "TASK-001",
                "reportType": "regular",
                "templateId": "regular-v1",
                "templateVersion": "1",
                "titleLines": ["测试证券营业部", "常规审计报告"],
                "addressee": "测试证券营业部：",
                "introduction": {"paragraphId": "intro", "text": "按照计划完成审计。", "evidenceIds": ["E-1"], "requiresHumanReview": False},
                "sections": [{
                    "heading": "一、审计项目基本情况",
                    "paragraphs": [{"paragraphId": "body", "text": "审计期间为2025年度。", "evidenceIds": ["E-1"], "requiresHumanReview": False}],
                    "tables": [],
                    "subsections": [],
                }],
                "closingOrganization": "测试证券公司",
                "reportDate": "2026年1月1日",
                "status": "ready-for-review",
                "blockers": [],
                "warnings": [],
                "allEvidenceIds": ["E-1"],
            }
            draft_path.write_text(json.dumps(draft, ensure_ascii=False), encoding="utf-8")

            subprocess.run(
                [sys.executable, str(script), "--template", str(template_path), "--draft", str(draft_path), "--output", str(output_path)],
                check=True,
            )
            rendered = Document(output_path)
            text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
            self.assertIn("测试证券营业部", text)
            self.assertIn("审计期间为2025年度。", text)
            self.assertIn("测试证券公司", text)
            self.assertEqual(rendered.paragraphs[0].runs[0].font.size, Pt(22))


if __name__ == "__main__":
    unittest.main()
