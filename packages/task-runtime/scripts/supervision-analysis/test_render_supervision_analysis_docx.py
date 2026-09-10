from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


class RenderSupervisionAnalysisDocxTest(unittest.TestCase):
    def test_renders_project_analysis_result_as_formal_docx(self) -> None:
        script = Path(__file__).with_name("render_supervision_analysis_docx.py")
        with tempfile.TemporaryDirectory() as raw_dir:
            root = Path(raw_dir)
            analysis_path = root / "analysis.json"
            records_path = root / "records.json"
            output_path = root / "report.docx"
            analysis = {
                "schemaVersion": "supervision-analysis.v1",
                "task": {"analysisStart": "2024-01-01", "analysisEnd": "2024-12-31"},
                "snapshot": {
                    "included": [
                        {
                            "documentId": "DOC-1",
                            "title": "监管决定测试资料",
                            "fileDate": "2024-01-02",
                            "organizationIds": ["东方证券股份有限公司"],
                            "dataOrigin": "official-public",
                            "uploadEntry": "file-center",
                            "sourceUrl": "https://example.test/source",
                        }
                    ]
                },
                "issues": [
                    {
                        "issueId": "ISSUE-1",
                        "sourceDocumentId": "DOC-1",
                        "organizationIds": ["东方证券股份有限公司"],
                        "category": "业务管理",
                        "description": "测试问题。",
                        "evidenceIds": ["VERSION-1:p1:issue-1"],
                    }
                ],
                "relations": [
                    {
                        "issueId": "ISSUE-1",
                        "relationType": "RECTIFICATION",
                        "status": "AUTO_CONFIRMED",
                        "selectedRecordId": "RECT-1",
                    }
                ],
                "statistics": {
                    "issueTotal": 1,
                    "pendingRelationReviewTotal": 0,
                    "byCategory": {"业务管理": 1},
                    "rectification": {
                        "confirmedTotal": 1,
                        "completionRate": 1.0,
                        "byStatus": {"COMPLETED": 1},
                    },
                    "accountability": {"confirmedTotal": 0},
                },
            }
            records = {
                "rectifications": [
                    {"recordId": "RECT-1", "status": "COMPLETED", "description": "已完成测试整改。"}
                ],
                "accountabilities": [],
            }
            analysis_path.write_text(json.dumps(analysis, ensure_ascii=False), encoding="utf-8")
            records_path.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
            subprocess.run(
                [
                    sys.executable,
                    str(script),
                    "--analysis",
                    str(analysis_path),
                    "--records",
                    str(records_path),
                    "--output",
                    str(output_path),
                    "--report-date",
                    "2026年9月2日",
                ],
                check=True,
            )
            document = Document(output_path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("监督共享信息汇总分析报告", text)
            self.assertIn("测试问题。", table_text)
            self.assertIn("已完成测试整改。", table_text)
            self.assertGreaterEqual(len(document.tables), 5)
            self.assertAlmostEqual(document.sections[0].page_width.cm, 21.0, places=1)
            self.assertAlmostEqual(document.sections[0].page_height.cm, 29.7, places=1)
            first_table = document.tables[0]
            self.assertEqual(first_table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")), "8640")
            self.assertIsNotNone(first_table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")))


if __name__ == "__main__":
    unittest.main()
