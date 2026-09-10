from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BODY_FONT = "宋体"
HEADING_FONT = "黑体"
ASCII_FONT = "Times New Roman"
GRAY_FILL = "E7E6E6"
LIGHT_GRAY_FILL = "F2F2F2"
WHITE = "FFFFFF"
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 89, 89)
TABLE_WIDTH_DXA = 8640
TABLE_INDENT_DXA = 120


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def set_run_font(run: Any, chinese: str = BODY_FONT, size: float = 12, bold: bool = False, color: RGBColor = BLACK) -> None:
    run.font.name = ASCII_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), ASCII_FONT)
    fonts.set(qn("w:hAnsi"), ASCII_FONT)
    fonts.set(qn("w:eastAsia"), chinese)


def set_paragraph_spacing(paragraph: Any, *, before: float = 0, after: float = 0, line: float = 1.5) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell: Any, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell: Any, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table: Any, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"Table widths must total {TABLE_WIDTH_DXA} DXA: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)


def repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_borders(table: Any, color: str = "7F7F7F", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def style_table_cell(cell: Any, text: str, *, header: bool = False, size: float = 9, center: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(paragraph, line=1.1)
    run = paragraph.add_run(text)
    set_run_font(run, chinese=HEADING_FONT if header else BODY_FONT, size=size, bold=header)
    if header:
        set_cell_shading(cell, GRAY_FILL)


def add_table(
    document: DocumentObject,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    font_size: float = 9,
    centered_columns: set[int] | None = None,
) -> Any:
    centered_columns = centered_columns or set()
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        style_table_cell(table.rows[0].cells[index], header, header=True, size=font_size, center=True)
    repeat_header(table.rows[0])
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            style_table_cell(row.cells[index], value, size=font_size, center=index in centered_columns)
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(document: DocumentObject, text: str, level: int) -> Any:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text)
    return paragraph


def add_body(document: DocumentObject, text: str, *, indent: bool = True, color: RGBColor = BLACK) -> Any:
    paragraph = document.add_paragraph(style="Normal")
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)
    run = paragraph.add_run(text)
    set_run_font(run, size=12, color=color)
    return paragraph


def add_note(document: DocumentObject, text: str) -> Any:
    paragraph = document.add_paragraph(style="Report Note")
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=MUTED)
    return paragraph


def add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)


def add_external_hyperlink(paragraph: Any, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0000FF")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), ASCII_FONT)
    fonts.set(qn("w:hAnsi"), ASCII_FONT)
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    run_properties.extend([fonts, color, underline, size])
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_document(document: DocumentObject) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    normal = document.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.5
    for level, size, before, after in ((1, 16, 14, 8), (2, 14, 10, 6), (3, 12, 8, 4)):
        style = document.styles[f"Heading {level}"]
        style.font.name = ASCII_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True
    note = document.styles.add_style("Report Note", 1)
    note.font.name = ASCII_FONT
    note.font.size = Pt(10.5)
    note.font.color.rgb = MUTED
    note._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    note.paragraph_format.left_indent = Pt(24)
    note.paragraph_format.right_indent = Pt(24)
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(6)
    note.paragraph_format.line_spacing = 1.25
    add_page_number(section.footer.paragraphs[0])


def add_title_block(document: DocumentObject, analysis: dict[str, Any], report_date: str) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(18)
    title_run = title.add_run("监督共享信息汇总分析报告")
    set_run_font(title_run, chinese=HEADING_FONT, size=22, bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run("（业务验收测试版）")
    set_run_font(subtitle_run, size=12)
    task = analysis["task"]
    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(4)
    metadata_run = metadata.add_run(f"分析期间：{task['analysisStart']}至{task['analysisEnd']}")
    set_run_font(metadata_run, size=10.5)
    generated = document.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated.paragraph_format.space_after = Pt(18)
    generated_run = generated.add_run(f"生成日期：{report_date}")
    set_run_font(generated_run, size=10.5)


def status_zh(status: str | None) -> str:
    return {
        "COMPLETED": "已完成",
        "IN_PROGRESS": "进行中",
        "PARTIALLY_COMPLETED": "部分完成",
        "CONTINUOUS": "持续整改",
        "AUTO_CONFIRMED": "自动确认",
        "HUMAN_CONFIRMED": "人工确认",
        "UNMATCHED": "待确认",
    }.get(status or "", status or "未提供")


def build_document(analysis: dict[str, Any], records: dict[str, Any], ocr: dict[str, Any] | None, report_date: str) -> DocumentObject:
    document = Document()
    configure_document(document)
    add_title_block(document, analysis, report_date)

    issues = analysis["issues"]
    relations = analysis["relations"]
    statistics = analysis["statistics"]
    materials = analysis["snapshot"]["included"]
    material_by_id = {item["documentId"]: item for item in materials}
    relation_by_key = {(item["issueId"], item["relationType"]): item for item in relations}
    rectifications = {item["recordId"]: item for item in records.get("rectifications", [])}
    raw_accountabilities = records.get("accountabilities", [])
    if isinstance(raw_accountabilities, dict):
        raw_accountabilities = [raw_accountabilities]
    accountabilities = {item["recordId"]: item for item in raw_accountabilities}

    add_heading(document, "一、摘要", 1)
    add_body(
        document,
        f"本次分析共纳入{len(materials)}份资料，其中公开监管资料6份、模拟内部资料2份。系统识别问题{statistics['issueTotal']}项，自动确认问题与整改关系{statistics['rectification']['confirmedTotal']}条、核验问责记录{statistics['accountability']['confirmedTotal']}条；另有{statistics['pendingRelationReviewTotal']}条关系尚未确认，保留待处理。",
    )
    add_body(
        document,
        "本报告中公开监管问题来源于证监会公开资料；整改、问责及内部权限问题为模拟测试数据，仅用于验证系统提取、关联、统计和报告生成能力，不代表东方证券真实内部整改或问责事实。",
    )
    category_rows = [[category, str(count), f"{count / statistics['issueTotal']:.1%}"] for category, count in statistics["byCategory"].items()]
    add_table(document, ["问题类别", "数量", "占比"], category_rows, [4500, 1800, 2340], font_size=10, centered_columns={1, 2})
    add_body(
        document,
        "制度及内控机制建设和业务管理各5项，合计占全部问题的71.4%。原始资料未提供统一风险等级字段，本报告不自行补充高、中、低风险判断。",
    )

    add_heading(document, "二、外部事项", 1)
    add_heading(document, "（一）监管检查、函件等指出问题及整改情况", 2)
    public_issues = [item for item in issues if not item["issueId"].startswith("SYN-")]
    grouped: dict[str, list[dict[str, Any]]] = {}
    for issue in public_issues:
        grouped.setdefault(issue["sourceDocumentId"], []).append(issue)
    for material_id, document_issues in sorted(grouped.items(), key=lambda item: material_by_id[item[0]].get("fileDate", "")):
        material = material_by_id[material_id]
        add_heading(document, material["title"], 3)
        source_summary = add_body(
            document,
            f"监管日期：{material.get('fileDate', '未提供')}。涉及主体：{'、'.join(material['organizationIds'])}。系统从该资料中识别{len(document_issues)}项问题。",
            indent=False,
        )
        # Keep each material summary with its following issue table so a page does
        # not end with a source title and summary while the table starts alone.
        source_summary.paragraph_format.keep_with_next = True
        rows: list[list[str]] = []
        for index, issue in enumerate(document_issues, 1):
            relation = relation_by_key[(issue["issueId"], "RECTIFICATION")]
            rectification_text = "待业务确认"
            selected_id = relation.get("selectedRecordId")
            if selected_id and selected_id in rectifications:
                record = rectifications[selected_id]
                rectification_text = f"{status_zh(record.get('status'))}；{record.get('description', '')}"
            rows.append([str(index), issue["category"], issue["description"], rectification_text, status_zh(relation["status"])])
        add_table(
            document,
            ["序号", "问题类别", "问题描述", "整改情况", "关联状态"],
            rows,
            [600, 1500, 2900, 2700, 940],
            font_size=8.5,
            centered_columns={0, 4},
        )
    add_note(document, "说明：上述整改数据均为模拟数据。“待业务确认”表示现有资料不足以建立精确关联，不表示相关问题未整改。")

    add_heading(document, "（二）外部审计、检查、调查等指出问题及整改情况", 2)
    add_body(document, "本次资料范围内未检索到外部审计报告、外部专项检查报告或外部调查报告，因此本模块不形成问题结论。该表述不等同于相关事项未发生。")

    add_heading(document, "三、内部事项", 1)
    synthetic_issues = [item for item in issues if item["issueId"].startswith("SYN-")]
    add_heading(document, "（一）内部审计、合规、风险检查发现问题及整改情况", 2)
    internal_rows: list[list[str]] = []
    for issue in synthetic_issues:
        relation = relation_by_key[(issue["issueId"], "RECTIFICATION")]
        selected_id = relation.get("selectedRecordId")
        record = rectifications.get(selected_id or "", {})
        internal_rows.append(
            [
                issue["issueId"],
                issue["description"],
                status_zh(record.get("status")),
                record.get("description", "未提供整改描述"),
            ]
        )
    add_table(document, ["问题编号", "问题描述", "整改状态", "整改措施"], internal_rows, [2200, 2600, 1200, 2640], font_size=9, centered_columns={2})
    add_note(document, "以上内部问题和整改记录均为模拟测试数据。")

    add_heading(document, "（二）问责处理情况", 2)
    accountability_rows: list[list[str]] = []
    for issue in synthetic_issues:
        relation = relation_by_key.get((issue["issueId"], "ACCOUNTABILITY"))
        if not relation:
            continue
        record = accountabilities.get(relation.get("selectedRecordId", ""), {})
        accountability_rows.append(
            [
                issue["issueId"],
                "、".join(record.get("responsibleDepartmentIds", [])) or "未提供",
                record.get("action", "未提供"),
                status_zh(relation["status"]),
            ]
        )
    if accountability_rows:
        add_table(document, ["关联问题", "责任部门", "处理措施", "关联状态"], accountability_rows, [2600, 2000, 2500, 1540], font_size=9, centered_columns={3})
    else:
        add_body(document, "本次资料范围内未形成已确认问责关联。")
    add_note(document, "以上问责记录为模拟测试数据。")

    add_heading(document, "（三）日常监督情况", 2)
    add_body(document, "当前资料未覆盖日常合规监督台账、风险监测台账或法律诉讼案件台账，本模块暂不形成统计结论。")

    add_heading(document, "四、整改及问责关联情况", 1)
    rectification = statistics["rectification"]
    status_rows = [[status_zh(key), str(value)] for key, value in rectification["byStatus"].items()]
    add_table(document, ["整改状态", "关联问题数"], status_rows, [4320, 4320], font_size=10, centered_columns={0, 1})
    add_body(
        document,
        f"按已确认且需要整改的问题总数计算，整改总体完成率为{rectification['completionRate']:.1%}。尚未关联整改资料的问题仍保留在分母中，资料缺失不等于整改未开展。由于当前整改数据为模拟数据，该比例不代表东方证券真实整改完成率。",
    )
    pending = [item for item in relations if item["status"] == "UNMATCHED"]
    if pending:
        pending_rows = []
        issue_by_id = {item["issueId"]: item for item in issues}
        for index, relation in enumerate(pending, 1):
            issue = issue_by_id[relation["issueId"]]
            pending_rows.append([str(index), issue["issueId"], issue["description"], "缺少可精确关联的问题编号或监管文号"])
        add_heading(document, "需业务确认事项", 2)
        add_table(document, ["序号", "问题编号", "问题摘要", "待确认原因"], pending_rows, [600, 2100, 3300, 2640], font_size=8.5, centered_columns={0})

    if ocr:
        add_heading(document, "五、资料处理情况", 1)
        page_total = sum(item["pageCount"] for item in ocr["documents"])
        block_total = sum(item["blockCount"] for item in ocr["documents"])
        chunk_total = sum(item["chunkCount"] for item in ocr["documents"])
        add_body(
            document,
            f"本批次共处理{len(ocr['documents'])}份、{page_total}页PDF资料，形成{block_total}个OCR内容块和{chunk_total}个检索块。全部资料状态为INDEXED，检索块已完成本地BGE-M3 dense和sparse向量生成及PG、Milvus投影。",
        )
        quality_rows = [
            [
                item["filename"],
                str(item["pageCount"]),
                str(item["blockCount"]),
                str(item["chunkCount"]),
                f"{item['ocrConfidenceMin']:.1%}",
                item["pipelineStatus"],
            ]
            for item in ocr["documents"]
        ]
        add_table(
            document,
            ["资料名称", "页数", "OCR块", "检索块", "最低置信度", "状态"],
            quality_rows,
            [3600, 600, 900, 900, 1300, 1340],
            font_size=8,
            centered_columns={1, 2, 3, 4, 5},
        )
        add_note(document, "本次未执行独立语义召回质量评测，当前结果证明OCR、向量生成和索引链路可运行，不代表召回准确率已经达到生产验收指标。")

    add_heading(document, "附件一：问题、整改、问责及证据总表", 1)
    detail_rows: list[list[str]] = []
    for index, issue in enumerate(issues, 1):
        material = material_by_id[issue["sourceDocumentId"]]
        rect_relation = relation_by_key[(issue["issueId"], "RECTIFICATION")]
        selected_rect = rectifications.get(rect_relation.get("selectedRecordId", ""), {})
        rect_text = "待业务确认" if rect_relation["status"] == "UNMATCHED" else f"{status_zh(selected_rect.get('status'))}；{selected_rect.get('description', '')}"
        acc_relation = relation_by_key.get((issue["issueId"], "ACCOUNTABILITY"))
        acc_text = "无已确认记录"
        if acc_relation and acc_relation.get("selectedRecordId"):
            acc_record = accountabilities.get(acc_relation["selectedRecordId"], {})
            acc_text = acc_record.get("action", "已确认")
        evidence = issue["evidenceIds"][0]
        page_number = evidence.split(":p", 1)[1].split(":", 1)[0] if ":p" in evidence else "未提供"
        detail_rows.append(
            [
                str(index),
                material.get("fileDate", "模拟资料"),
                "、".join(issue["organizationIds"]),
                issue["category"],
                issue["description"],
                rect_text,
                acc_text,
                f"第{page_number}页",
            ]
        )
    add_table(
        document,
        ["序号", "日期", "涉及主体", "问题类别", "问题描述", "整改情况", "问责", "证据页"],
        detail_rows,
        [480, 900, 1300, 1000, 1800, 1700, 900, 560],
        font_size=7.3,
        centered_columns={0, 1, 7},
    )

    add_heading(document, "附件二：资料清单", 1)
    source_rows: list[list[str]] = []
    for index, material in enumerate(materials, 1):
        source_rows.append(
            [
                str(index),
                material["title"],
                "证监会公开资料" if material["dataOrigin"] == "official-public" else "模拟内部资料",
                material.get("fileDate", "模拟"),
                "文件中心" if material["uploadEntry"] == "file-center" else "监督共享",
            ]
        )
    source_table = add_table(document, ["序号", "资料名称", "来源性质", "日期", "处理入口"], source_rows, [500, 4400, 1400, 1100, 1240], font_size=8.5, centered_columns={0, 3, 4})
    for row_index, material in enumerate(materials, 1):
        if not material.get("sourceUrl"):
            continue
        cell = source_table.rows[row_index].cells[1]
        paragraph = cell.add_paragraph()
        set_paragraph_spacing(paragraph, line=1.0)
        add_external_hyperlink(paragraph, "公开来源", material["sourceUrl"])

    return document


def main() -> None:
    parser = argparse.ArgumentParser(description="Render a supervision-analysis.v1 result as a formal DOCX report")
    parser.add_argument("--analysis", required=True, type=Path, help="supervision-analysis.v1 JSON")
    parser.add_argument("--records", required=True, type=Path, help="JSON containing rectifications/accountabilities")
    parser.add_argument("--ocr", type=Path, help="optional OCR batch export JSON")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report-date", help="Chinese report date; defaults to current local date")
    args = parser.parse_args()
    analysis = read_json(args.analysis)
    if analysis.get("schemaVersion") != "supervision-analysis.v1":
        raise ValueError("--analysis must contain schemaVersion=supervision-analysis.v1")
    records = read_json(args.records)
    ocr = read_json(args.ocr) if args.ocr else None
    report_date = args.report_date or f"{date.today().year}年{date.today().month}月{date.today().day}日"
    document = build_document(analysis, records, ocr, report_date)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(json.dumps({"output": str(args.output), "schemaVersion": analysis["schemaVersion"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
